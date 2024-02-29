import os
import streamlit as st
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
from jira import JIRA, JIRAError
from langchain.text_splitter import CharacterTextSplitter
from langchain.embeddings.openai import OpenAIEmbeddings
from langchain import FAISS
from langchain.chains.question_answering import load_qa_chain
from langchain.llms import OpenAI
from langchain.callbacks import get_openai_callback
import base64

API_KEY = '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'
embeddings_global = OpenAIEmbeddings(openai_api_key=API_KEY)

os.environ["JIRA_API_TOKEN"] = "ATATT3xFfGF0RffC8ROCQtuzGolQ10pwSK4YCrX_JDRlFYjCMoYbF46Cdtkv71k20YqJrXIjtt5EWg3LM2cATzq5okY27HgRO1NbE_sl89eFlvox0fz9aZbBnZhnog1n_CFnAKmhhuZkEVLnjfQuAiQJUwbqfOYBKZXjDPn611vtrWG48ySlRys=2C3FA0A6"
os.environ["JIRA_USERNAME"] = "Swapnil Gangwal"
os.environ["JIRA_INSTANCE_URL"] = "https://ai-crafter.atlassian.net/jira/software/projects/HACKT/pages"
os.environ["OPENAI_API_KEY"] = "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"

# JIRA connection details (Replace with your actual details)
jira_options = {'server': 'https://ai-crafter.atlassian.net/jira/software/projects/HACKT/pages'}
jira_user = 'Swapnil Gangwal' #swpnl.08@gmail.com
jira_api_token = 'ATATT3xFfGF0RffC8ROCQtuzGolQ10pwSK4YCrX_JDRlFYjCMoYbF46Cdtkv71k20YqJrXIjtt5EWg3LM2cATzq5okY27HgRO1NbE_sl89eFlvox0fz9aZbBnZhnog1n_CFnAKmhhuZkEVLnjfQuAiQJUwbqfOYBKZXjDPn611vtrWG48ySlRys=2C3FA0A6'
jira_project_key = 'PROJECT_KEY'
jira_issue_type = 'Task'


# Example usage
image_path = 'ChatBot1.png'


def create_jira_ticket(summary, description):
  try:
      jira = JIRA(options=jira_options, basic_auth=(jira_user, jira_api_token))
      issue_dict = {
          'project': {'key': jira_project_key},
          'summary': summary,
          'description': description,
          'issuetype': {'name': jira_issue_type},
      }
      new_issue = jira.create_issue(fields=issue_dict)
      return new_issue.key
  except JIRAError as e:
      print(f"Failed to create JIRA ticket: {e}")
      return None
  except Exception as e:
      # This catches other unforeseen exceptions
      print(f"An unexpected error occurred: {e}")
      return None
      
def process_text(text):
    text_splitter = CharacterTextSplitter(
        separator="\n",
        chunk_size=1000,
        chunk_overlap=200,
        length_function=len
    )
    chunks = text_splitter.split_text(text)
    embeddings = OpenAIEmbeddings(openai_api_key=API_KEY)
    knowledgeBase = FAISS.from_texts(chunks, embeddings)
    return knowledgeBase

# Functions to read PDF, DOCX, and XLSX files remain the same
def read_pdf(file):
    pdf_reader = PdfReader(file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text() or ""  # Fallback to empty string if None
    return text

def read_docx(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def read_xlsx(file):
    wb = openpyxl.load_workbook(file)
    text = ""
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            if row:
                text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
    return text
def main():
    # CSS to inject contained in a multi-line string
    background = """
        <style>
        h1 {
            color: #ADD8E6;
        }
        .stTextInput>div>div>input {
            color: black;
        }
        .stButton>button {
            border: 2px solid #ADD8E6; /* Light Blue */
            background-color: #ADD8E6; /* Light Blue */
            color: white;
            padding: 10px 24px;
            cursor: pointer;
            border-radius: 5px;
        }
        .stButton>button:hover {
            background-color: #BFEFFF; /* Lighter Blue */
            color: white;
        }
        </style>
    """


    st.markdown(background, unsafe_allow_html=True)
    st.markdown("# AI Crafters' 💬")
    st.image(image_path, width=700)

    uploaded_files = st.file_uploader('Upload your document(s)', type=['pdf', 'docx', 'xlsx'], accept_multiple_files=True)

    text = ""
    if uploaded_files:
        for uploaded_file in uploaded_files:
            if uploaded_file.type == "application/pdf":
                text += read_pdf(uploaded_file)
            elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                text += read_docx(uploaded_file)
            elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                text += read_xlsx(uploaded_file)
        
        knowledgeBase = process_text(text)

        query = st.text_input('Ask a question to the document(s)')
        cancel_button = st.button('Cancel')

        if cancel_button:
            st.stop()
        no_answer = ["i'm sorry", "i don't know", "i don't know what","I do not know,+"]
        if query:
            docs = knowledgeBase.similarity_search(query)
            llm = OpenAI(openai_api_key=API_KEY)
            chain = load_qa_chain(llm, chain_type='stuff')

            with get_openai_callback() as cost:
                response = chain.run(input_documents=docs, question=query)

            if not response or 'no answer found' in response.lower() or any(substring in response.lower() for substring in no_answer):
              st.write("No answer found. Would you like to raise a JIRA ticket?")
              col1, col2 = st.columns(2)  # Create two columns for Yes and No buttons
              if col1.button('Yes'):
                  ticket_id = create_jira_ticket(
                      summary=f"Query unanswered: {query[:50]}...",  # Truncate query to fit summary
                      description=f"A query made to documents did not return an answer: {query}\n\nPlease investigate."
                  )
                  st.success(f"JIRA ticket created successfully: {ticket_id}")
              elif col2.button('No'):
                  st.info("You chose not to raise a JIRA ticket.")
            else:
                st.write(response)



if __name__ == "__main__":
    main()
